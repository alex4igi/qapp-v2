#!/usr/bin/env python3
"""Extrage calendarele metodologice din xlsx într-un JSON stabil, revizuibil.

De ce Python și nu o librărie npm: SheetJS 0.18.5 (singura versiune de pe npm)
strică perechile surrogate — 🌴 U+1F334 ajunge U+F334 — și are două CVE-uri high
fără fix pe registry. openpyxl citește fișierele corect.

JSON-ul rezultat e artefactul durabil: importul (import-programe-metodologice.mjs)
citește DOAR din el, deci re-rularea importului nu depinde de Python sau de xlsx.
Rulează scriptul ăsta doar când se schimbă fișierele Excel.

Rulare:  python3 scripts/migrate/extract-calendar-metodologic.py
Necesită: pip install openpyxl
"""

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SURSA = ROOT / "Calendar metodologic teacheri"
IESIRE = Path(__file__).resolve().parent / "data" / "calendar-metodologic-2026-2027.json"
SHEETS_IGNORATE = {"Sumar grupe", "Ghid activități"}


def celula(rand, i):
    v = rand[i] if i < len(rand) else None
    return str(v).strip() if v is not None else ""


def parseaza_sheet(ws):
    randuri = list(ws.iter_rows(values_only=True))
    eticheta = celula(randuri[0], 0) if randuri else ""
    meta = [p.strip() for p in celula(randuri[1], 0).split("·")] if len(randuri) > 1 else []

    def dupa(prefix):
        for p in meta:
            if p.startswith(prefix):
                return p[len(prefix):].strip()
        return ""

    rezultat = {
        "eticheta": eticheta,
        "locatie": meta[0] if meta else "",
        "studio": meta[1] if len(meta) > 1 else "",
        "ora": meta[2] if len(meta) > 2 else "",
        "nivel": dupa("Nivel:"),
        "zile": dupa("Zile:"),
        "module": [],
        "vacante": [],
        "lectii": [],
    }

    modul_curent = None
    for rand in randuri:
        a = celula(rand, 0)
        if not a:
            continue

        m = re.match(r"^MODUL\s+(\d+)\s*\((.+?)\)\s*$", a)
        if m:
            modul_curent = {
                "numar": int(m.group(1)),
                "nume": f"MODUL {m.group(1)}",
                "interval": m.group(2).strip(),
                "tema": None,
                "subtitlu": None,
            }
            rezultat["module"].append(modul_curent)
            continue

        m = re.match(r"^(?:🌴\s*)?VACANȚĂ\s+(\d+)\s*\((.+?)\)\s*[—–-]+\s*(.*)$", a)
        if m:
            rezultat["vacante"].append({
                "numar": int(m.group(1)),
                "nume": f"VACANȚĂ {m.group(1)}",
                "interval": m.group(2).strip(),
                "nota": m.group(3).strip() or None,
            })
            continue

        if a.startswith("▸") and modul_curent is not None:
            subtitlu = a.lstrip("▸").strip()
            modul_curent["subtitlu"] = subtitlu
            bucati = subtitlu.split("·")
            modul_curent["tema"] = bucati[-1].strip() if len(bucati) > 1 else None
            continue

        m = re.match(r"^Ședința\s+(\d+)$", a)
        if m and modul_curent is not None:
            titlu = celula(rand, 1)
            rezultat["lectii"].append({
                "nr_sedinta": int(m.group(1)),
                "modul": modul_curent["numar"],
                "titlu": titlu,
                "note": celula(rand, 2) or None,
            })

    return rezultat


def main():
    grupe = []
    for fisier in sorted(SURSA.glob("Calendar_Lectii_*.xlsx")):
        wb = openpyxl.load_workbook(fisier, read_only=True)
        for nume in wb.sheetnames:
            if nume in SHEETS_IGNORATE:
                continue
            grupa = parseaza_sheet(wb[nume])
            grupa["fisier"] = fisier.name
            grupa["sheet"] = nume
            grupe.append(grupa)
        wb.close()

    IESIRE.parent.mkdir(parents=True, exist_ok=True)
    IESIRE.write_text(json.dumps({"sezon": "2026-2027", "grupe": grupe}, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"{len(grupe)} grupe -> {IESIRE.relative_to(ROOT)}")
    for g in grupe:
        print(f"  {g['sheet']:22} {len(g['lectii']):3} lecții, {len(g['module'])} module, {len(g['vacante'])} vacanțe")


if __name__ == "__main__":
    main()
